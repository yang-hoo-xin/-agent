from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SUMMARY_DIR = ROOT / "tmp" / "summaries"
OUTPUT = ROOT / "output" / "documents" / "大模型应用算法面试题与八股复习手册.docx"

PARTS = [
    ("第一部分", "Part 1 第 1 至 185 页", SUMMARY_DIR / "part1_1_185.md"),
    ("第二部分", "Part 1 第 186 至 371 页", SUMMARY_DIR / "part1_186_371.md"),
    ("第三部分", "Part 2 第 1 至 162 页", SUMMARY_DIR / "part2_1_162.md"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_font(run, name="Microsoft YaHei", size=10.5, bold=None, color="000000") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)
    set_font(run, size=9, color="666666")


def clean_inline(text: str) -> str:
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    return text.strip()


def clean_heading(text: str) -> str:
    text = clean_inline(text)
    text = re.sub(r"^\d+(?:\.\d+)*[\.、]?\s*", "", text)
    text = text.replace("–", " 至 ").replace("—", " 至 ").replace("-", " 至 ")
    text = re.sub(r"[《》【】（）()：:；;，,、。！？!?/\\|]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.18

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(16)

    for name, size, before, after in (
        ("Heading 1", 16, 16, 7),
        ("Heading 2", 13, 12, 5),
        ("Heading 3", 11.5, 9, 4),
    ):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    question = styles.add_style("Interview Question", 1)
    question.font.name = "Microsoft YaHei"
    question._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    question.font.size = Pt(10.8)
    question.font.bold = True
    question.font.color.rgb = RGBColor.from_string("17365D")
    question.paragraph_format.space_before = Pt(7)
    question.paragraph_format.space_after = Pt(2)
    question.paragraph_format.keep_with_next = True

    answer = styles.add_style("Interview Answer", 1)
    answer.font.name = "Microsoft YaHei"
    answer._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    answer.font.size = Pt(10.5)
    answer.font.color.rgb = RGBColor(0, 0, 0)
    answer.paragraph_format.left_indent = Cm(0.35)
    answer.paragraph_format.space_after = Pt(5)
    answer.paragraph_format.line_spacing = 1.18


def add_intro(doc: Document) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("大模型应用算法面试题与八股复习手册")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("基于两份学习资料整理的完整复习版")
    set_font(r, size=12, color="404040")
    p.paragraph_format.space_after = Pt(24)

    p = doc.add_paragraph()
    r = p.add_run(
        "本手册汇总 533 页材料中的面试题和八股知识，共整理 318 组问答。内容覆盖 RAG、Agent、"
        "MCP、记忆与多 Agent、评估与可观测性、微调与推理、工程系统设计、行为题和算法题。"
        "各章保留来源页码，便于回到原材料核对上下文。"
    )
    set_font(r, size=11)
    p.paragraph_format.space_after = Pt(14)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(6.1), Cm(4.1), Cm(4.1)]
    headers = ["来源", "页数", "问答数量"]
    for i, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "17365D")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_font(r, size=10, bold=True, color="FFFFFF")
    rows = [
        ("Part 1 第 1 至 185 页", "185", "114"),
        ("Part 1 第 186 至 371 页", "186", "119"),
        ("Part 2 第 1 至 162 页", "162", "85"),
    ]
    for row_idx, row_data in enumerate(rows, start=1):
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, row_data)):
            cell.width = widths[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx % 2 == 0:
                set_cell_shading(cell, "EEF3F8")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_font(r, size=9.8)
    set_table_borders(table)

    doc.add_paragraph("使用方法", style="Heading 1")
    items = [
        "第一遍按主题阅读八股要点，先建立知识结构。",
        "第二遍只看问题，先口头回答，再对照答案补充遗漏。",
        "第三遍按来源页码回查难题，把真实项目数据和个人经历替换进答案。",
        "版本、价格、榜单和发布时间变化较快，正式面试前需要再次核实。",
    ]
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)

    doc.add_paragraph("主题目录", style="Heading 1")
    topic_rows = [
        (
            "第一部分  Part 1 第 1 至 185 页",
            "学习与岗位路线 RAG 项目真题 Agent 项目 LLM 与 Transformer 模型选型 RAG 全链路 Agent 框架与评估",
        ),
        (
            "第二部分  Part 1 第 186 至 371 页",
            "MCP 记忆 多 Agent A2A 幻觉 OpenClaw Hermes Loop 与 Graph Pi Agent 工程容错 微调推理 Harness Skill Claude Code",
        ),
        (
            "第三部分  Part 2 第 1 至 162 页",
            "Claude Code 记忆与压缩 Agent Team 权限与成本 Vibe Coding RAG 评估上线 微调推理 Python 与 LangGraph 系统设计 行为题 算法题",
        ),
    ]
    for title, topics in topic_rows:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(title)
        set_font(r, size=10.5, bold=True)
        q = doc.add_paragraph(style="List Bullet 2")
        q.add_run(topics)


def add_markdown_part(doc: Document, part_name: str, label: str, path: Path) -> None:
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(part_name)
    q = doc.add_paragraph()
    r = q.add_run(label)
    set_font(r, size=12, bold=True, color="404040")
    q.paragraph_format.space_after = Pt(18)

    lines = path.read_text(encoding="utf-8").splitlines()
    for line in lines:
        raw = line.strip()
        if not raw:
            continue
        if raw.startswith("# ") or raw.startswith(">"):
            continue
        if raw.startswith("## "):
            doc.add_paragraph(clean_heading(raw[3:]), style="Heading 1")
            continue
        if raw.startswith("### "):
            doc.add_paragraph(clean_heading(raw[4:]), style="Heading 2")
            continue
        if raw.startswith("#### "):
            doc.add_paragraph(clean_heading(raw[5:]), style="Heading 3")
            continue

        text = clean_inline(raw)
        if re.match(r"^Q\s*[：:]", text, flags=re.I):
            p = doc.add_paragraph(style="Interview Question")
            p.add_run(text)
            continue
        if re.match(r"^A\s*[：:]", text, flags=re.I):
            p = doc.add_paragraph(style="Interview Answer")
            p.add_run(text)
            continue
        if raw.startswith(("- ", "• ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(clean_inline(raw[2:]))
            continue
        if re.match(r"^\d+[\.、]\s*", raw):
            p = doc.add_paragraph(style="List Number")
            p.add_run(clean_inline(re.sub(r"^\d+[\.、]\s*", "", raw)))
            continue
        p = doc.add_paragraph()
        p.add_run(text)


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.1)
        section.right_margin = Cm(2.1)
        section.footer_distance = Cm(0.9)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_field(footer)

    add_intro(doc)
    for part_name, label, path in PARTS:
        add_markdown_part(doc, part_name, label, path)

    doc.core_properties.title = "大模型应用算法面试题与八股复习手册"
    doc.core_properties.subject = "大模型应用算法面试复习"
    doc.core_properties.keywords = "大模型 RAG Agent MCP 微调 推理 面试"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
